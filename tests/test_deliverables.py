"""Tests for the deliverable files (Excel, Word, PowerPoint) and the Manager's
read-back checks, for both test months.

Run from the project root with:   python -m unittest -v
"""

import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

import config
from agents.action_log import ActionLog
from agents.manager import create_manager
from agents.simulated_ai import ManagerPolicy
from tools import excel_report, slides_report
from tools.csv_helpers import read_csv
from tools.deliverable_checks import check_excel, check_slides, check_word
from tools.manager_tools import ManagerDesk
from tools.review_checks import ReportReview
from tools.writer_tools import numbers_in


def run_company(month, output_dir, **options):
    config.set_month(month)
    with redirect_stdout(io.StringIO()):
        manager = create_manager(ActionLog(output_dir / "logs"), output_dir=output_dir, **options)
        manager.run(f"Prepare the {config.month_name()} report.")
    return manager


class DeliveredFilesTests(unittest.TestCase):
    """Runs both months once and checks the delivered files."""

    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.runs = {}
        for month in ["2026-09", "2026-10"]:
            output_dir = Path(cls.temp_dir.name) / month
            output_dir.mkdir()
            manager = run_company(month, output_dir)
            cls.runs[month] = {
                "dir": output_dir, "manager": manager,
                "analysis": json.loads((output_dir / config.ANALYSIS_FILE_NAME).read_text(encoding="utf-8")),
                "texts": json.loads((output_dir / config.DRAFT_TEXTS_FILE_NAME).read_text(encoding="utf-8")),
                "categorized": read_csv(output_dir / config.CATEGORIZED_FILE_NAME),
                "exceptions": read_csv(output_dir / config.EXCEPTIONS_FILE_NAME),
                "files": {kind: output_dir / name for kind, name in config.REPORT_DELIVERABLES.items()},
            }
        config.set_month(config.DEFAULT_MONTH)

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()
        config.set_month(config.DEFAULT_MONTH)

    def each_month(self):
        return list(self.runs.items())

    def test_files_are_delivered_with_the_report(self):
        for month, run in self.each_month():
            self.assertEqual(run["manager"].desk.finished, "delivered")
            for kind, path in run["files"].items():
                self.assertTrue(path.exists(), path.name)
                draft = run["dir"] / f"draft_report_{month}.{kind}"
                self.assertEqual(path.read_bytes(), draft.read_bytes())

    def test_read_back_checks_find_no_problems(self):
        for month, run in self.each_month():
            allowed = set(numbers_in(json.dumps(run["analysis"], ensure_ascii=False)))
            self.assertEqual(check_excel(run["files"]["xlsx"], run["analysis"], run["categorized"],
                                         run["exceptions"]), [])
            self.assertEqual(check_word(run["files"]["docx"], run["analysis"], run["exceptions"],
                                        run["texts"], allowed), [])
            self.assertEqual(check_slides(run["files"]["pptx"], run["analysis"], allowed), [])

    def test_review_notes_mention_the_files(self):
        for month, run in self.each_month():
            notes = (run["dir"] / f"review_notes_{month}.md").read_text(encoding="utf-8")
            self.assertIn("[PASS] Deliverables match the Analyst's results", notes)

    # --- Excel ------------------------------------------------------------------

    def test_excel_layout_and_formats(self):
        for month, run in self.each_month():
            workbook = load_workbook(run["files"]["xlsx"])
            self.assertEqual(workbook.sheetnames, ["Summary", "Transactions", "Exceptions"])
            for sheet in workbook.worksheets:
                self.assertEqual(sheet.freeze_panes, "A2", sheet.title)          # frozen header row
                self.assertGreater(sheet.column_dimensions["C"].width, 10, sheet.title)
            summary = workbook["Summary"]
            self.assertEqual(summary["B2"].number_format, excel_report.CZK)
            self.assertEqual(summary["D2"].number_format, excel_report.CZK_SIGNED)
            self.assertEqual(summary["E2"].number_format, excel_report.PERCENT_SIGNED)
            self.assertEqual(workbook["Transactions"]["E2"].number_format, excel_report.CZK)

    def test_excel_variance_columns_are_formulas(self):
        for month, run in self.each_month():
            summary = load_workbook(run["files"]["xlsx"])["Summary"]
            self.assertEqual(summary["D2"].value, "=B2-C2")
            self.assertEqual(summary["E2"].value, '=IF(C2=0,"",D2/C2)')
            self.assertEqual(summary["A9"].value, "Total expenses")
            self.assertEqual(summary["B9"].value, "=SUM(B2:B8)")
            self.assertEqual(summary["D9"].value, "=B9-C9")

    def test_excel_actuals_are_sumifs_over_transactions(self):
        for month, run in self.each_month():
            summary = load_workbook(run["files"]["xlsx"])["Summary"]
            # Whole columns: a row added to Transactions later is still counted.
            self.assertEqual(summary["B2"].value, "=-SUMIFS(Transactions!$E:$E,Transactions!$F:$F,$A2)")
            self.assertEqual(summary["B11"].value, '=SUMIFS(Transactions!$E:$E,Transactions!$F:$F,"Sales")')
            self.assertIsInstance(summary["C2"].value, int)          # the budget stays an input

    def test_excel_print_setup(self):
        for month, run in self.each_month():
            for sheet in load_workbook(run["files"]["xlsx"]).worksheets:
                self.assertEqual(sheet.page_setup.orientation, "landscape", sheet.title)
                self.assertEqual(sheet.page_setup.fitToWidth, 1, sheet.title)
                self.assertEqual(sheet.page_setup.fitToHeight, 0, sheet.title)
                self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage, sheet.title)
                self.assertEqual(sheet.print_title_rows, "$1:$1", sheet.title)   # header row on every page

    def test_no_double_negatives_in_any_format(self):
        from tools.deliverable_checks import slides_text, word_text
        from tools.writer_tools import sign_wording_problems
        for month, run in self.each_month():
            report = (run["dir"] / f"report_{month}.md").read_text(encoding="utf-8")
            for text in [report, word_text(run["files"]["docx"]), slides_text(run["files"]["pptx"])]:
                self.assertEqual(sign_wording_problems(text), [])
            self.assertIn(run["analysis"]["budget_vs_actual"]["total"]["variance_phrase"],
                          word_text(run["files"]["docx"]))

    def test_powerpoint_action_titles(self):
        expected = {
            "2026-09": ["September 3.0% over budget; Travel and Marketing overspent, Software may have missing costs",
                        "Salaries and Rent close to budget; biggest variance: Marketing, CZK 17,000 (68.0%) over budget",
                        "Overspend in Travel and Marketing; Software more than 20% under budget",
                        "Three actions before closing September"],
            "2026-10": ["October 5.5% under budget; three categories may have missing costs",
                        "Salaries and Rent close to budget; biggest variance: Travel, CZK 11,470 (76.5%) under budget",
                        "Travel, Software and Marketing more than 20% under budget – check for missing or late costs",
                        "Four actions before closing October"],
        }
        for month, run in self.each_month():
            deck = Presentation(run["files"]["pptx"])
            titles = [slide.shapes[0].text_frame.text.replace("\u00a0", " ") for slide in deck.slides]
            self.assertEqual(titles, expected[month])

    def test_powerpoint_headline_has_cards_and_to_do_line_only(self):
        for month, run in self.each_month():
            first = " ".join(shape.text_frame.text for shape in Presentation(run["files"]["pptx"]).slides[0].shapes
                             if shape.has_text_frame)
            self.assertIn("to clear – see the last slide", first)
            for item in run["analysis"]["budget_vs_actual"]["categories"]:
                self.assertNotIn(item["category"] + ":", first)     # category detail is on slide 3

    def test_powerpoint_mentions_the_refund_in_its_category(self):
        run = self.runs["2026-10"]
        self.assertEqual(run["analysis"]["totals"]["refunds"], {"Travel": {"total_czk": "3,100", "lines": [25]}})
        self.assertEqual(self.runs["2026-09"]["analysis"]["totals"]["refunds"], {})
        refund = {"total_czk": "3,100", "lines": [25]}
        self.assertEqual(slides_report.refund_note(refund, "over"), "; after a refund of CZK 3,100 (line 25)")
        self.assertEqual(slides_report.refund_note(None, "under"), "")
        unusual = " ".join(shape.text_frame.text for shape in Presentation(run["files"]["pptx"]).slides[2].shapes
                           if shape.has_text_frame).replace("\u00a0", " ")
        self.assertIn("Travel: actual 3,530 vs budget 15,000, CZK 11,470 (76.5%) under budget; "
                      "CZK 3,100 of the gap is a refund (line 25)", unusual)
        september = " ".join(shape.text_frame.text for shape in
                             Presentation(self.runs["2026-09"]["files"]["pptx"]).slides[2].shapes
                             if shape.has_text_frame)
        self.assertNotIn("refund", september)                       # no refunds in September

    def test_powerpoint_closing_panels_start_under_the_title(self):
        for month, run in self.each_month():
            closing = Presentation(run["files"]["pptx"]).slides[3]
            boxes = [shape for shape in closing.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
            self.assertTrue(boxes)
            self.assertTrue(all(abs(shape.top - Inches(1.6)) < Inches(0.01) for shape in boxes))

    def test_powerpoint_has_no_explanations_or_repeated_paragraphs(self):
        for month, run in self.each_month():
            deck = Presentation(run["files"]["pptx"])
            all_text = " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes
                                if shape.has_text_frame)
            self.assertNotIn("separate scales", all_text)
            closing = " ".join(shape.text_frame.text for shape in deck.slides[3].shapes if shape.has_text_frame)
            self.assertNotIn(run["texts"]["exceptions_comment"], closing)
            self.assertNotIn("Before the month is closed", closing)

    def test_word_columns_never_split_a_word(self):
        from docx.shared import Cm
        for month, run in self.each_month():
            for table in Document(run["files"]["docx"]).tables:
                for column_index, column in enumerate(table.columns):
                    words = [w for cell in column.cells for w in cell.text.split()]
                    longest = max((len(w) for w in words), default=0)
                    # Word stores widths in twips, so allow a rounding difference of 0.01 cm.
                    self.assertGreaterEqual(column.width, Cm((longest + 2) * 0.2) - Cm(0.01),
                                            f"{month}: column {column_index} narrower than its longest word")

    def test_excel_transactions_and_exceptions_sheets(self):
        for month, run in self.each_month():
            workbook = load_workbook(run["files"]["xlsx"])
            transactions = list(workbook["Transactions"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(transactions), len(run["categorized"]))
            headers = [cell.value for cell in workbook["Transactions"][1]]
            self.assertIn("Needs review", headers)
            self.assertIn("Review reason", headers)
            exception_lines = {row[0] for row in workbook["Exceptions"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(exception_lines, {int(r["source_line"]) for r in run["exceptions"]})

    # --- Word and PowerPoint ----------------------------------------------------

    def test_word_has_the_report_content_and_tables(self):
        for month, run in self.each_month():
            document = Document(run["files"]["docx"])
            text = "\n".join(p.text for p in document.paragraphs)
            title = "September 2026" if month == "2026-09" else "October 2026"
            self.assertIn(f"Monthly expense report – {title}", text)
            self.assertIn(run["texts"]["summary"], text)
            self.assertIn("Probable duplicates – confirm against the bank statement", text)
            self.assertGreaterEqual(len(document.tables), 2)
            self.assertEqual([c.text for c in document.tables[0].rows[0].cells],
                             ["Category", "Actual", "Budget", "Variance", "Variance %"])

    def test_powerpoint_has_at_most_five_slides_and_two_charts(self):
        for month, run in self.each_month():
            deck = Presentation(run["files"]["pptx"])
            self.assertLessEqual(len(deck.slides), 5)
            charts = [shape for slide in deck.slides for shape in slide.shapes if shape.has_chart]
            self.assertEqual(len(charts), 2)   # large categories and the rest, separate scales


class TamperedFileTests(unittest.TestCase):
    """A file changed after it was built (or a generator bug) must block delivery."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_dir = Path(self.temp_dir.name)
        config.set_month("2026-10")
        self.log = ActionLog(self.output_dir / "logs")
        with redirect_stdout(io.StringIO()):
            self.desk = ManagerDesk(self.log, output_dir=self.output_dir)
            for step in [self.desk.ask_collector, self.desk.ask_categorizer, self.desk.ask_analyst,
                         self.desk.ask_writer]:
                step()
        self.files = {kind: self.output_dir / name for kind, name in config.DRAFT_DELIVERABLES.items()}

    def tearDown(self):
        self.temp_dir.cleanup()
        config.set_month(config.DEFAULT_MONTH)

    def review(self):
        result = ReportReview(output_dir=self.output_dir).run()
        failed = {c["check"]: c for c in result["checks"] if not c["passed"]}
        issues = [i for i in result["issues"] if i["check"] == "Deliverables match the Analyst's results"]
        return failed, issues

    def assert_blocked(self, text):
        failed, issues = self.review()
        self.assertIn("Deliverables match the Analyst's results", failed)
        self.assertTrue(all(i["owner"] == "data" for i in issues))
        self.assertTrue(any(text in i["issue"] for i in issues), [i["issue"] for i in issues])
        with redirect_stdout(io.StringIO()):
            self.desk.review_report()
        with self.assertRaises(RuntimeError):
            self.desk.deliver_report("note")

    def test_untouched_files_pass(self):
        failed, _ = self.review()
        self.assertEqual(failed, {})

    def test_typed_number_instead_of_sumifs_is_caught(self):
        workbook = load_workbook(self.files["xlsx"])
        workbook["Summary"]["B4"] = 3530          # Travel actual typed in: right value, but no longer traceable
        workbook.save(self.files["xlsx"])
        self.assert_blocked("Excel Summary row 4 (Travel): actual is not the formula =-SUMIFS(")

    def test_wrong_amount_in_transactions_sheet_is_caught(self):
        workbook = load_workbook(self.files["xlsx"])
        sheet = workbook["Transactions"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == 25:                # the Berlin hotel refund, +3,100
                row[4].value = 3010
        workbook.save(self.files["xlsx"])
        self.assert_blocked("Excel Summary row 4 (Travel): actual works out to 3620, the Analyst says 3,530")

    def test_row_added_below_the_transactions_is_counted(self):
        # The SUMIFS cover whole columns, so the check works the added row in too.
        workbook = load_workbook(self.files["xlsx"])
        sheet = workbook["Transactions"]
        sheet.append([None, None, "Added by hand", "", -5000, "Travel"])
        workbook.save(self.files["xlsx"])
        self.assert_blocked("Excel Summary row 4 (Travel): actual works out to 8530, the Analyst says 3,530")

    def test_formula_replaced_by_a_number_is_caught(self):
        workbook = load_workbook(self.files["xlsx"])
        workbook["Summary"]["D4"] = -11470        # right value, but no longer a formula
        workbook.save(self.files["xlsx"])
        self.assert_blocked("variance is not the formula =B4-C4")

    def test_wrong_number_in_word_is_caught(self):
        document = Document(self.files["docx"])
        document.tables[0].rows[3].cells[1].paragraphs[0].runs[0].text = "3,350"
        document.save(self.files["docx"])
        self.assert_blocked("Word budget table, Travel")

    def test_wrong_number_in_powerpoint_chart_is_caught(self):
        deck = Presentation(self.files["pptx"])
        chart = [s for s in deck.slides[1].shapes if s.has_chart][1].chart
        data = CategoryChartData()
        data.categories = list(chart.plots[0].categories)
        values = {s.name: list(s.values) for s in chart.plots[0].series}
        values["Actual"][0] = 3350.0             # Travel
        data.add_series("Actual", values["Actual"])
        data.add_series("Budget", values["Budget"])
        chart.replace_data(data)
        deck.save(self.files["pptx"])
        self.assert_blocked("PowerPoint chart: actual/budget differ")

    def test_missing_file_goes_back_to_the_writer(self):
        self.files["pptx"].unlink()
        failed, _ = self.review()
        self.assertIn("Deliverables are present and up to date", failed)

    def test_file_changed_after_review_cannot_be_delivered(self):
        with redirect_stdout(io.StringIO()):
            self.assertTrue(self.desk.review_report()["passed"])
        with open(self.files["docx"], "ab") as file:
            file.write(b" ")
        with self.assertRaises(RuntimeError):
            self.desk.deliver_report("note")


class WordingCheckTests(unittest.TestCase):
    def test_writer_rejects_double_negatives(self):
        from tools.writer_tools import ReportWriter
        writer = ReportWriter()
        writer.analysis = {"total": "-33,514", "percent": "-5.5%"}
        with self.assertRaises(ValueError) as caught:
            writer.save_report("Total expenses were CZK -33,514 under budget.", "Nothing to explain.")
        self.assertIn("Double negative", str(caught.exception))

    def test_manager_catches_a_double_negative_in_a_file(self):
        with tempfile.TemporaryDirectory() as folder:
            output_dir = Path(folder)
            run_company("2026-10", output_dir)
            draft = output_dir / config.DRAFT_DELIVERABLES["docx"]
            for name in config.REPORT_DELIVERABLES.values():
                (output_dir / name).unlink()
            document = Document(draft)
            document.add_paragraph("Travel was -11,470 under budget.")
            document.save(draft)
            result = ReportReview(output_dir=output_dir).run()
        failed = {c["check"]: c for c in result["checks"] if not c["passed"]}
        self.assertIn("No signed amounts with over/under budget", failed)
        issues = [i for i in result["issues"] if i["check"] == "No signed amounts with over/under budget"]
        self.assertEqual([i["owner"] for i in issues], ["writer"])
        self.assertIn('Double negative in the docx: "-11,470 under budget"', issues[0]["issue"])
        config.set_month(config.DEFAULT_MONTH)


class GeneratorBugTests(unittest.TestCase):
    def tearDown(self):
        config.set_month(config.DEFAULT_MONTH)

    def test_manager_escalates_when_a_generator_writes_a_wrong_number(self):
        real_write_summary = excel_report.write_summary

        def buggy_write_summary(sheet, analysis, month_label):
            real_write_summary(sheet, analysis, month_label)
            sheet["C2"] = 405000                  # Salaries budget off by 45,000

        with tempfile.TemporaryDirectory() as folder, \
                mock.patch.object(excel_report, "write_summary", buggy_write_summary):
            output_dir = Path(folder)
            manager = run_company("2026-10", output_dir, fake_script=ManagerPolicy())
            escalation = (output_dir / config.ESCALATION_FILE_NAME).read_text(encoding="utf-8")
            delivered = [name for name in config.REPORT_DELIVERABLES.values() if (output_dir / name).exists()]

        self.assertEqual(manager.desk.finished, "escalated")
        self.assertEqual(delivered, [])
        self.assertIn("budget 405000 differs from the Analyst's 450,000", escalation)


if __name__ == "__main__":
    unittest.main()
