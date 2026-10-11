"""The Excel deliverable: report_YYYY-MM.xlsx.

Sheets:
  Summary       budget vs actual per category. Actual is a SUMIFS over the
                Transactions sheet, so every total traces back to its rows;
                budget is an input; variance and variance % are formulas.
  Transactions  every categorized row, with its review flag and reason
  Exceptions    every row set aside (duplicates, outside the period,
                missing amounts, other issues)

Everything is written by Python from the Analyst's results and the
categorized data - the AI never types a number into this file.
"""

from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
CZK = '#,##0 "CZK";-#,##0 "CZK";"-"'
CZK_SIGNED = '+#,##0 "CZK";-#,##0 "CZK";"-"'
PERCENT_SIGNED = '+0.0%;-0.0%;0.0%'
DATE = "yyyy-mm-dd"

HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1D3352")
BODY_FONT = Font(name=FONT)
INPUT_FONT = Font(name=FONT, color="0000FF")          # blue: inputs (the budget)
LINK_FONT = Font(name=FONT, color="008000")           # green: formulas reading another sheet
BOLD_FONT = Font(name=FONT, bold=True)
NOTE_FONT = Font(name=FONT, italic=True, color="555555")
TOTAL_BORDER = Border(top=Side(style="thin"), bottom=Side(style="double"))
TOP = Alignment(vertical="top")   # rows with wrapped text stay aligned

SUMMARY_HEADERS = ["Category", "Actual (CZK)", "Budget (CZK)", "Variance (CZK)", "Variance %"]
TRANSACTION_HEADERS = ["Line", "Date", "Description", "Counterparty", "Amount (CZK)", "Category",
                       "Method", "Reason", "Needs review", "Review reason"]
EXCEPTION_HEADERS = ["Line", "Date", "Description", "Counterparty", "Amount (CZK)", "Issue", "Detail"]


def money(text):
    """'42,000' or '+10,546' -> Decimal."""
    return Decimal(text.replace(",", "").replace("+", ""))


def as_number(amount):
    """Decimal -> int when whole (clean in Excel), else float."""
    return int(amount) if amount == amount.to_integral_value() else float(amount)


def as_date(text):
    try:
        return date.fromisoformat(text)
    except ValueError:
        return text   # an unreadable date stays as it was


def actual_formula(row, last_transaction_row):
    """Actual spend of the category in column A: minus the sum of its rows in
    Transactions (expenses are negative there, refunds positive)."""
    return (f"=-SUMIFS(Transactions!$E$2:$E${last_transaction_row},"
            f"Transactions!$F$2:$F${last_transaction_row},$A{row})")


def sales_formula(last_transaction_row):
    """Sales income: the sum of the Sales rows in Transactions (positive amounts)."""
    return (f'=SUMIFS(Transactions!$E$2:$E${last_transaction_row},'
            f'Transactions!$F$2:$F${last_transaction_row},"Sales")')


def last_transaction_row(categorized_rows):
    return max(2, len(categorized_rows) + 1)


def variance_formula(row):
    return f"=B{row}-C{row}"


def percent_formula(row):
    return f'=IF(C{row}=0,"",D{row}/C{row})'


def write_header(sheet, headers):
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"   # the header row stays visible when scrolling
    set_up_printing(sheet)


def set_up_printing(sheet):
    """Landscape, one page wide (as many pages tall as needed), header row on every page."""
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"


def set_widths(sheet, minimum=8, maximum=36):
    """Column width from the longest value in each column. Long text columns
    are capped and wrap instead, so the sheet still prints one page wide."""
    for column_cells in sheet.columns:
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = max(minimum, min(maximum, longest + 3))
        if longest + 3 > maximum:
            for cell in column_cells[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(analysis, categorized_rows, exception_rows, path, month_label):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    write_summary(summary, analysis, month_label, last_transaction_row(categorized_rows))
    write_transactions(workbook.create_sheet("Transactions"), categorized_rows)
    write_exceptions(workbook.create_sheet("Exceptions"), exception_rows)
    workbook.save(path)


def write_summary(sheet, analysis, month_label, last_row):
    write_header(sheet, SUMMARY_HEADERS)
    categories = analysis["budget_vs_actual"]["categories"]
    first = 2
    last = first + len(categories) - 1

    for row, item in enumerate(categories, start=first):
        sheet.cell(row=row, column=1, value=item["category"]).font = BODY_FONT
        cell = sheet.cell(row=row, column=2, value=actual_formula(row, last_row))
        cell.font, cell.number_format = LINK_FONT, CZK
        cell = sheet.cell(row=row, column=3, value=as_number(money(item["budget_czk"])))
        cell.font, cell.number_format = INPUT_FONT, CZK
        cell = sheet.cell(row=row, column=4, value=variance_formula(row))
        cell.font, cell.number_format = BODY_FONT, CZK_SIGNED
        cell = sheet.cell(row=row, column=5, value=percent_formula(row))
        cell.font, cell.number_format = BODY_FONT, PERCENT_SIGNED

    total = last + 1
    sheet.cell(row=total, column=1, value="Total expenses")
    sheet.cell(row=total, column=2, value=f"=SUM(B{first}:B{last})").number_format = CZK
    sheet.cell(row=total, column=3, value=f"=SUM(C{first}:C{last})").number_format = CZK
    sheet.cell(row=total, column=4, value=variance_formula(total)).number_format = CZK_SIGNED
    sheet.cell(row=total, column=5, value=percent_formula(total)).number_format = PERCENT_SIGNED
    for column in range(1, 6):
        sheet.cell(row=total, column=column).font = BOLD_FONT
        sheet.cell(row=total, column=column).border = TOTAL_BORDER

    sales_row = total + 2
    sheet.cell(row=sales_row, column=1, value="Sales income (side note)").font = BODY_FONT
    cell = sheet.cell(row=sales_row, column=2, value=sales_formula(last_row))
    cell.font, cell.number_format = LINK_FONT, CZK

    notes = [f"{month_label}. Actual = minus the sum of the category's rows in Transactions (expenses are "
             f"negative there, refunds positive). Variance = Actual - Budget; + means over budget. "
             f"Variance % = Variance / Budget.",
             "Blue = budget (input); green = formulas reading the Transactions sheet; black = formulas."]
    understated = analysis["exceptions"]["understated_categories"]
    if understated:
        parts = "; ".join(f"{category} (line {', '.join(str(n) for n in lines)})"
                          for category, lines in understated.items())
        notes.append(f"Actual understated - amount still missing for: {parts}. See the Exceptions sheet.")
    for offset, note in enumerate(notes, start=2):
        # Each note is wrapped inside the table's width (columns A-E), so it
        # doesn't run off to the right and widen the printed page.
        row = sales_row + offset
        cell = sheet.cell(row=row, column=1, value=note)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.row_dimensions[row].height = 13 * (len(note) // 95 + 1)

    sheet.column_dimensions["A"].width = 26
    for letter in "BCDE":
        sheet.column_dimensions[letter].width = 18


def write_transactions(sheet, rows):
    write_header(sheet, TRANSACTION_HEADERS)
    for index, row in enumerate(rows, start=2):
        values = [int(row["source_line"]), as_date(row["date"]), row["description"], row["counterparty"],
                  as_number(Decimal(row["amount_czk"])), row["category"], row["method"], row["reason"],
                  row.get("needs_review", "no"), row.get("review_reason", "")]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.font = BODY_FONT
            cell.alignment = TOP
        sheet.cell(row=index, column=2).number_format = DATE
        sheet.cell(row=index, column=5).number_format = CZK
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TRANSACTION_HEADERS))}{max(2, len(rows) + 1)}"
    set_widths(sheet)


def write_exceptions(sheet, rows):
    write_header(sheet, EXCEPTION_HEADERS)
    for index, row in enumerate(rows, start=2):
        amount = as_number(Decimal(row["amount_czk"])) if row["amount_czk"].strip() else None
        values = [int(row["source_line"]), as_date(row["date"]), row["description"], row["counterparty"],
                  amount, row["issue"], row["detail"]]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.font = BODY_FONT
            cell.alignment = TOP
        sheet.cell(row=index, column=2).number_format = DATE
        sheet.cell(row=index, column=5).number_format = CZK
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXCEPTION_HEADERS))}{len(rows) + 1}"
    else:
        sheet.cell(row=2, column=1, value="No rows were set aside this month.").font = NOTE_FONT
    set_widths(sheet)
