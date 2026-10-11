"""The PowerPoint deliverable: report_YYYY-MM.pptx (4 slides).

  1. Headline numbers
  2. Budget vs actual - two bar charts: the large categories (Salaries,
     Rent) dwarf the rest, so they get their own chart and axis instead of
     squashing the others flat
  3. Unusual items
  4. Before closing the month - exceptions and to-do

Every number comes from the Analyst's results (or the Writer's checked text).
"""

from decimal import Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1D, 0x33, 0x52)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xE8, 0xED, 0xF3)
ACTUAL_COLOR = RGBColor(0x2A, 0x78, 0xD6)   # categorical slot 1 (blue)
BUDGET_COLOR = RGBColor(0xEB, 0x68, 0x34)   # categorical slot 2 (orange) - validated pair
FONT = "Calibri"

# Categories with a budget at least this large go into the "large categories" chart.
LARGE_BUDGET = Decimal(50000)


def count(number, singular, plural):
    return f"{number} {singular if number == 1 else plural}"


def lines_label(items):
    numbers = [str(x["source_line"]) for x in items]
    return ("line " if len(numbers) == 1 else "lines ") + ", ".join(numbers)


def money(text):
    return Decimal(text.replace(",", "").replace("+", ""))


def text_box(slide, left, top, width, height, paragraphs, size=16, color=None):
    """paragraphs: list of strings or (text, bold) tuples, one per line."""
    frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    frame.word_wrap = True
    for index, item in enumerate(paragraphs):
        text, bold = item if isinstance(item, tuple) else (item, False)
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
        run.font.color.rgb = color or NAVY
        paragraph.space_after = Pt(6)
    return frame


def title(slide, text, subtitle=None):
    text_box(slide, 0.6, 0.35, 12.1, 0.8, [(text, True)], size=28)
    if subtitle:
        text_box(slide, 0.6, 1.05, 12.1, 0.5, [subtitle], size=14, color=GREY)


def kpi(slide, left, label, value, note=""):
    box = slide.shapes.add_shape(1, Inches(left), Inches(2.2), Inches(2.85), Inches(2.1))  # rectangle
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.fill.background()
    text_box(slide, left + 0.15, 2.35, 2.55, 0.5, [label], size=14, color=GREY)
    text_box(slide, left + 0.15, 2.85, 2.55, 0.8, [(value, True)], size=24)
    if note:
        text_box(slide, left + 0.15, 3.6, 2.55, 0.6, [note], size=12, color=GREY)


def bar_chart(slide, left, width, heading, rows):
    text_box(slide, left, 1.6, width, 0.5, [(heading, True)], size=16)
    data = CategoryChartData(number_format='#,##0')
    data.categories = [row["category"] for row in rows]
    data.add_series("Actual", [float(money(row["actual_czk"])) for row in rows])
    data.add_series("Budget", [float(money(row["budget_czk"])) for row in rows])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(left), Inches(2.1),
                                   Inches(width), Inches(4.6), data)
    chart = shape.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    for series, color in zip(chart.series, [ACTUAL_COLOR, BUDGET_COLOR]):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    plot.data_labels.number_format = '#,##0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.font.color.rgb = GREY
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.category_axis.reverse_order = True            # first category at the top
    value_axis = chart.value_axis
    value_axis.tick_labels.font.size = Pt(10)
    value_axis.tick_labels.number_format = '#,##0'
    value_axis.tick_labels.number_format_is_linked = False
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = LIGHT
    return chart


def build_presentation(analysis, texts, path, month_label):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)   # 16:9
    blank = deck.slide_layouts[6]

    headline_slide(deck.slides.add_slide(blank), analysis, month_label)
    chart_slide(deck.slides.add_slide(blank), analysis)
    unusual_slide(deck.slides.add_slide(blank), analysis, texts)
    closing_slide(deck.slides.add_slide(blank), analysis, texts)
    deck.save(path)


def headline_slide(slide, analysis, month_label):
    total = analysis["budget_vs_actual"]["total"]
    totals = analysis["totals"]
    title(slide, f"Monthly expense report – {month_label}", "LedgerMind · all amounts in CZK")
    kpi(slide, 0.6, "Total expenses", total["actual_czk"])
    kpi(slide, 3.65, "Budget", total["budget_czk"])
    kpi(slide, 6.7, "Variance", total["variance_czk"], f"{total['variance_percent']} vs budget")
    kpi(slide, 9.75, "Sales income", totals["sales_income_czk"], "side note")
    unusual = analysis["unusual_items"]
    over = ", ".join(x["category"] for x in unusual["categories_over_budget"]) or "none"
    under = ", ".join(x["category"] for x in unusual["categories_under_budget"]) or "none"
    text_box(slide, 0.6, 4.8, 12.1, 1.6,
             [f"More than {unusual['over_budget_threshold']} over budget: {over}",
              f"More than {unusual['over_budget_threshold']} under budget: {under}"], size=16)


def chart_slide(slide, analysis):
    rows = analysis["budget_vs_actual"]["categories"]
    large = [row for row in rows if money(row["budget_czk"]) >= LARGE_BUDGET]
    small = [row for row in rows if money(row["budget_czk"]) < LARGE_BUDGET]
    title(slide, "Budget vs actual", "Two charts with separate scales: the large categories would "
                                     "otherwise squash the rest")
    if large:
        bar_chart(slide, 0.6, 4.6, "Large categories", large)
    if small:
        bar_chart(slide, 5.5 if large else 0.6, 7.2 if large else 12.1, "Other categories", small)


def unusual_slide(slide, analysis, texts):
    unusual = analysis["unusual_items"]
    threshold = unusual["over_budget_threshold"]
    lines = [(f"More than {threshold} over budget", True)]
    lines += [f"{x['category']}: actual {x['actual_czk']} vs budget {x['budget_czk']} "
              f"({x['variance_czk']}, {x['variance_percent']})" for x in unusual["categories_over_budget"]] or ["None"]
    lines += [(f"More than {threshold} under budget", True)]
    lines += [f"{x['category']}: actual {x['actual_czk']} vs budget {x['budget_czk']} "
              f"({x['variance_czk']}, {x['variance_percent']}) – {x['note']}"
              for x in unusual["categories_under_budget"]] or ["None"]
    lines += [(f"Single expenses above {unusual['large_transaction_threshold_czk']} "
               f"(excluding {' and '.join(unusual['large_transactions_excluding'])})", True)]
    lines += [f"Line {x['source_line']}: {x['description']} ({x['category']}) – {x['amount_czk']}"
              for x in unusual["large_transactions"]] or ["None"]
    title(slide, "Unusual items")
    text_box(slide, 0.6, 1.4, 12.1, 3.6, lines, size=15)
    if texts.get("unusual_items_comment"):
        text_box(slide, 0.6, 5.2, 12.1, 1.9, [texts["unusual_items_comment"]], size=13, color=GREY)


def closing_slide(slide, analysis, texts):
    exceptions = analysis["exceptions"]

    todo = []
    if exceptions["probable_duplicates_count"]:
        todo.append(f"Confirm {count(exceptions['probable_duplicates_count'], 'probable duplicate', 'probable duplicates')} "
                    f"against the bank statement ({lines_label(exceptions['probable_duplicates'])}; "
                    f"total {exceptions['probable_duplicates_total_czk']})")
    if exceptions["possible_duplicates_count"]:
        todo.append(f"Confirm {count(exceptions['possible_duplicates_count'], 'possible duplicate', 'possible duplicates')} "
                    f"against the bank statement ({lines_label(exceptions['possible_duplicates'])}; "
                    f"total {exceptions['possible_duplicates_total_czk']})")
    if exceptions["outside_the_period_count"]:
        todo.append(f"Book {count(exceptions['outside_the_period_count'], 'row', 'rows')} from another period "
                    f"in its own month ({lines_label(exceptions['outside_the_period'])})")
    if exceptions["missing_amounts_count"]:
        todo.append(f"Obtain {count(exceptions['missing_amounts_count'], 'missing amount', 'missing amounts')} "
                    f"({lines_label(exceptions['missing_amounts'])}) – the categories they belong to are understated")
    for item in exceptions["needs_review"]:
        todo.append(f"Review line {item['source_line']} {item['description']} ({item['category']}, "
                    f"{item['amount_czk']})")
    if not todo:
        todo = ["No exceptions this month"]

    title(slide, "Before closing the month", "Excluded from the totals until confirmed, or included but to be checked")
    text_box(slide, 0.6, 1.6, 12.1, 4.2, ["• " + line for line in todo], size=15)
    if texts.get("exceptions_comment"):
        text_box(slide, 0.6, 6.0, 12.1, 1.2, [texts["exceptions_comment"]], size=13, color=GREY)
