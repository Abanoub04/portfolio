"""The Word deliverable: report_YYYY-MM.docx.

Rendered from the same content blocks as the Markdown report
(tools/report_content.py), so it always says exactly the same - formatted
as a professional document with real Word tables.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
TEXT_WIDTH_CM = 17.0          # A4 width minus 2 cm margins on both sides
NAVY = RGBColor(0x1D, 0x33, 0x52)
HEADER_SHADING = "1D3352"
TOTAL_SHADING = "E8EDF3"


def shade(cell, color):
    """Background colour of a table cell."""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def repeat_as_header(row):
    """Repeat the table's header row at the top of every page."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_up_document():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)          # A4
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 22), ("Heading 1", 15)]:
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY
        style.font.bold = True
    return document


def add_runs(paragraph, runs):
    for text, bold in runs:
        paragraph.add_run(text).bold = bold


def add_table(document, block):
    headers, rows = block["headers"], block["rows"]
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for column, title in enumerate(headers):
        cell = table.rows[0].cells[column]
        cell.text = ""
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if block["align"][column] == "right":
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        shade(cell, HEADER_SHADING)
    repeat_as_header(table.rows[0])

    for index, values in enumerate(rows, start=1):
        is_total = block["bold_last_row"] and index == len(rows)
        for column, value in enumerate(values):
            cell = table.rows[index].cells[column]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            run.bold = is_total
            if block["align"][column] == "right":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if is_total:
                shade(cell, TOTAL_SHADING)

    set_column_widths(table, headers, rows)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph()   # space after the table


def set_column_widths(table, headers, rows):
    """Short columns (line, date, amounts) get the width of their longest
    value so they never wrap; long text columns share the rest of the page."""
    cm_per_character = 0.2          # roughly, at 9 pt
    natural = []
    for column, title in enumerate(headers):
        longest = max([len(title)] + [len(row[column]) for row in rows])
        natural.append((longest + 2) * cm_per_character)
    short = [w <= 3.0 for w in natural]
    fixed = sum(w for w, is_short in zip(natural, short) if is_short)
    flexible = sum(w for w, is_short in zip(natural, short) if not is_short)
    room = TEXT_WIDTH_CM - fixed
    widths = [w if is_short else (w * room / flexible if flexible else w) for w, is_short in zip(natural, short)]
    if not flexible:   # only short columns: stretch them all to the page width
        widths = [w * TEXT_WIDTH_CM / fixed for w in widths]

    table.autofit = False
    for column, width_cm in enumerate(widths):
        width = Cm(width_cm)
        table.columns[column].width = width      # the table grid (used by LibreOffice)
        for row in table.rows:
            row.cells[column].width = width      # each cell (used by Word)


def build_document(blocks, path):
    document = set_up_document()
    for block in blocks:
        if block["type"] == "heading":
            document.add_heading(block["text"], level=0 if block["level"] == 1 else 1)
        elif block["type"] == "paragraph":
            add_runs(document.add_paragraph(), block["runs"])
        elif block["type"] == "bullets":
            for item in block["items"]:
                add_runs(document.add_paragraph(style="List Bullet"), item)
        elif block["type"] == "table":
            add_table(document, block)
    document.save(path)
